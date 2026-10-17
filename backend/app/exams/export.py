"""Export services for Exam results in Excel (.xlsx) and PDF formats."""

from __future__ import annotations

import html
import io
from datetime import UTC, datetime, timezone
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.models import Exam


def sanitize_excel_value(val: Any) -> Any:
    """
    Sanitize values to prevent CSV / Spreadsheet formula injection.
    Formulas starting with =, +, -, @, \t, or \r are prefixed with a single quote.
    """
    if val is None:
        return ""
    if isinstance(val, (int, float, bool)):
        return val
    s = str(val)
    if s and s[0] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + s
    return s


def sanitize_text(val: Any) -> str:
    """Sanitize arbitrary text for PDF paragraph rendering."""
    if val is None:
        return ""
    s = str(val)
    return html.escape(s)


def generate_exam_results_excel(exam: Exam, results: list[dict]) -> bytes:
    """Generate professional Excel workbook of exam results with sanitized cells."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Exam Results"

    # Color palette
    header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    meta_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    font_title = Font(name="Calibri", size=14, bold=True, color="1E293B")
    font_meta_label = Font(name="Calibri", size=10, bold=True, color="475569")
    font_meta_val = Font(name="Calibri", size=10, color="1E293B")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10, color="0F172A")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )

    # 1. Metadata Header Section
    ws.append([sanitize_excel_value(f"Exam Results: {exam.title or 'Untitled Exam'}")])
    ws.cell(row=1, column=1).font = font_title
    ws.merge_cells("A1:H1")

    exam_subject = exam.subject.value if hasattr(exam.subject, "value") else str(exam.subject or "N/A")
    metadata_rows = [
        ("Subject:", exam_subject, "Max Marks:", exam.maxMarks or 0),
        (
            "Exported On:",
            datetime.now(UTC).strftime("%Y-%m-%d %H:%M:%S UTC"),
            "Total Attempts:",
            len(results),
        ),
    ]

    for label1, val1, label2, val2 in metadata_rows:
        ws.append([label1, sanitize_excel_value(val1), "", label2, sanitize_excel_value(val2)])
        row_idx = ws.max_row
        ws.cell(row=row_idx, column=1).font = font_meta_label
        ws.cell(row=row_idx, column=2).font = font_meta_val
        ws.cell(row=row_idx, column=4).font = font_meta_label
        ws.cell(row=row_idx, column=5).font = font_meta_val

    ws.append([])  # Spacer row

    # 2. Table Headers
    headers = [
        "Rank",
        "Student Name",
        "Roll No",
        "Marks Obtained",
        "Max Marks",
        "Percentage (%)",
        "Status",
        "Submitted At",
    ]
    ws.append(headers)
    table_header_row = ws.max_row

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=table_header_row, column=col_idx)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = cell_border

    # 3. Data Rows
    if not results:
        ws.append(["No student attempts recorded for this exam."])
        ws.merge_cells(f"A{ws.max_row}:H{ws.max_row}")
        ws.cell(row=ws.max_row, column=1).alignment = Alignment(horizontal="center")
        ws.cell(row=ws.max_row, column=1).font = Font(italic=True, color="64748B")
    else:
        for i, item in enumerate(results):
            sub_at = item.get("submittedAt")
            sub_at_str = (
                sub_at.strftime("%Y-%m-%d %H:%M")
                if isinstance(sub_at, datetime)
                else (str(sub_at) if sub_at else "Not submitted")
            )
            row_data = [
                sanitize_excel_value(item.get("rank", i + 1)),
                sanitize_excel_value(item.get("studentName", "Unknown")),
                sanitize_excel_value(item.get("rollNo") or "-"),
                sanitize_excel_value(item.get("marksObtained", 0)),
                sanitize_excel_value(item.get("maxMarks", exam.maxMarks or 0)),
                sanitize_excel_value(item.get("percentage", 0.0)),
                sanitize_excel_value(item.get("status", "UNKNOWN")),
                sanitize_excel_value(sub_at_str),
            ]
            ws.append(row_data)
            current_row = ws.max_row

            is_alt = (i % 2 == 1)
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.font = font_data
                cell.border = cell_border
                if is_alt:
                    cell.fill = alt_fill

                # Alignments
                if col_idx in (1, 4, 5, 6):
                    cell.alignment = Alignment(horizontal="right")
                elif col_idx in (3, 7, 8):
                    cell.alignment = Alignment(horizontal="center")
                else:
                    cell.alignment = Alignment(horizontal="left")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_exam_results_pdf(exam: Exam, results: list[dict]) -> bytes:
    """Generate printable PDF report of exam results."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(letter),
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ExamTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#0F172A"),
        alignment=0,
    )
    meta_style = ParagraphStyle(
        "ExamMeta",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#475569"),
    )
    th_style = ParagraphStyle(
        "TH",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9,
        leading=11,
        textColor=colors.white,
        alignment=1,
    )
    td_style = ParagraphStyle(
        "TD",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#1E293B"),
    )
    td_center_style = ParagraphStyle(
        "TDCenter",
        parent=td_style,
        alignment=1,
    )
    td_right_style = ParagraphStyle(
        "TDRight",
        parent=td_style,
        alignment=2,
    )

    elements = []

    # Title & Subtitle
    clean_title = sanitize_text(exam.title or "Untitled Exam")
    elements.append(Paragraph(f"Exam Results: {clean_title}", title_style))
    elements.append(Spacer(1, 8))

    exam_subject = exam.subject.value if hasattr(exam.subject, "value") else str(exam.subject or "N/A")
    meta_text = (
        f"<b>Subject:</b> {sanitize_text(exam_subject)} &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"<b>Max Marks:</b> {exam.maxMarks or 0} &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"<b>Total Candidates:</b> {len(results)} &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"<b>Generated:</b> {datetime.now(UTC).strftime('%Y-%m-%d %H:%M UTC')}"
    )
    elements.append(Paragraph(meta_text, meta_style))
    elements.append(Spacer(1, 14))

    # Table
    headers = [
        Paragraph("Rank", th_style),
        Paragraph("Student Name", th_style),
        Paragraph("Roll No", th_style),
        Paragraph("Marks", th_style),
        Paragraph("Max", th_style),
        Paragraph("Percentage", th_style),
        Paragraph("Status", th_style),
        Paragraph("Submitted At", th_style),
    ]

    table_data = [headers]

    if not results:
        table_data.append(
            [
                Paragraph("<i>No student attempts recorded for this exam.</i>", td_center_style),
                "",
                "",
                "",
                "",
                "",
                "",
                "",
            ]
        )
    else:
        for i, item in enumerate(results):
            sub_at = item.get("submittedAt")
            sub_at_str = (
                sub_at.strftime("%Y-%m-%d %H:%M")
                if isinstance(sub_at, datetime)
                else (str(sub_at) if sub_at else "Not submitted")
            )
            pct = item.get("percentage", 0.0)
            pct_str = f"{pct:.1f}%" if isinstance(pct, (int, float)) else str(pct)

            row = [
                Paragraph(str(item.get("rank", i + 1)), td_center_style),
                Paragraph(sanitize_text(item.get("studentName", "Unknown")), td_style),
                Paragraph(sanitize_text(item.get("rollNo") or "-"), td_center_style),
                Paragraph(str(item.get("marksObtained", 0)), td_right_style),
                Paragraph(str(item.get("maxMarks", exam.maxMarks or 0)), td_right_style),
                Paragraph(pct_str, td_right_style),
                Paragraph(sanitize_text(item.get("status", "UNKNOWN")), td_center_style),
                Paragraph(sanitize_text(sub_at_str), td_center_style),
            ]
            table_data.append(row)

    # Column widths for landscape letter (792 pt width, margins 36*2=72, printable ~720 pt)
    col_widths = [40, 160, 70, 50, 50, 70, 80, 100]
    results_table = Table(table_data, colWidths=col_widths, repeatRows=1)

    t_style = TableStyle(
        [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#CBD5E1")),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]
    )

    if not results:
        t_style.add("SPAN", (0, 1), (-1, 1))
    else:
        for idx in range(1, len(table_data)):
            if idx % 2 == 0:
                t_style.add("BACKGROUND", (0, idx), (-1, idx), colors.HexColor("#F8FAFC"))

    results_table.setStyle(t_style)
    elements.append(results_table)

    doc.build(elements)
    return buf.getvalue()
