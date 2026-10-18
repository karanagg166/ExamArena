"""Unit tests for Exam Results Excel and PDF export functions and sanitization."""

import io
from datetime import UTC, datetime, timezone
from unittest.mock import AsyncMock, MagicMock

import openpyxl
import pytest
from reportlab.pdfgen import canvas

from app.core.models import Exam, Subject
from app.exams.export import (
    generate_exam_results_excel,
    generate_exam_results_pdf,
    sanitize_excel_value,
    sanitize_text,
)


def test_sanitize_excel_value_formula_injection():
    # Dangerous formula prefixes
    assert sanitize_excel_value("=SUM(A1:A10)") == "'=SUM(A1:A10)"
    assert sanitize_excel_value("+123456789") == "'+123456789"
    assert sanitize_excel_value("-50") == "'-50"
    assert sanitize_excel_value("@ADMIN") == "'@ADMIN"
    assert sanitize_excel_value("\tcmd") == "'\tcmd"
    assert sanitize_excel_value("\rcmd") == "'\rcmd"

    # Benign strings
    assert sanitize_excel_value("Alice Smith") == "Alice Smith"
    assert sanitize_excel_value("ROLL-102") == "ROLL-102"

    # Numeric and None values
    assert sanitize_excel_value(100) == 100
    assert sanitize_excel_value(95.5) == 95.5
    assert sanitize_excel_value(True) is True
    assert sanitize_excel_value(None) == ""


def test_sanitize_text_xml_escaping():
    assert sanitize_text("Alice <script>alert(1)</script>") == "Alice &lt;script&gt;alert(1)&lt;/script&gt;"
    assert sanitize_text("Marks & Percentage") == "Marks &amp; Percentage"
    assert sanitize_text(None) == ""


def test_generate_exam_results_excel_with_injection_and_unicode():
    exam = MagicMock(spec=Exam)
    exam.id = "ex_101"
    exam.title = "Physics Midterm & Final <2026>"
    exam.subject = Subject.SCIENCE
    exam.maxMarks = 100

    results = [
        {
            "rank": 1,
            "attemptId": "att_1",
            "studentId": "std_1",
            "studentName": "=CMD|' /C calc'!A0",  # Malicious formula injection attempt
            "rollNo": "+9199999999",             # Phone/formula injection
            "marksObtained": 95,
            "maxMarks": 100,
            "percentage": 95.0,
            "status": "SUBMITTED",
            "startedAt": datetime(2026, 10, 15, 10, 0, tzinfo=UTC),
            "submittedAt": datetime(2026, 10, 15, 11, 30, tzinfo=UTC),
        },
        {
            "rank": 2,
            "attemptId": "att_2",
            "studentId": "std_2",
            "studentName": "राहुल शर्मा (Rahul Sharma) 🚀",  # Unicode & Emoji
            "rollNo": "R-102",
            "marksObtained": 80,
            "maxMarks": 100,
            "percentage": 80.0,
            "status": "SUBMITTED",
            "startedAt": datetime(2026, 10, 15, 10, 5, tzinfo=UTC),
            "submittedAt": datetime(2026, 10, 15, 11, 45, tzinfo=UTC),
        },
        {
            "rank": 3,
            "attemptId": "att_3",
            "studentId": "std_3",
            "studentName": "Bob unsubmitted",
            "rollNo": None,                       # Missing roll number
            "marksObtained": 0,
            "maxMarks": 100,
            "percentage": 0.0,
            "status": "IN_PROGRESS",              # In progress / unsubmitted
            "startedAt": datetime(2026, 10, 15, 10, 10, tzinfo=UTC),
            "submittedAt": None,                  # Missing submission date
        },
    ]

    xlsx_bytes = generate_exam_results_excel(exam, results)
    assert xlsx_bytes.startswith(b"PK")  # ZIP / XLSX magic bytes

    # Verify workbook parsing and cell content
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active
    assert "Exam Results" in ws.title

    # Row 1: Title, Rows 2-3: Meta, Row 4: Spacer, Row 5: Headers, Row 6: First student
    student_1_name_val = ws.cell(row=6, column=2).value
    student_1_roll_val = ws.cell(row=6, column=3).value
    assert student_1_name_val == "'=CMD|' /C calc'!A0"
    assert student_1_roll_val == "'+9199999999"

    # Check unicode name
    student_2_name_val = ws.cell(row=7, column=2).value
    assert "राहुल शर्मा" in student_2_name_val

    # Check unsubmitted student row
    student_3_roll_val = ws.cell(row=8, column=3).value
    student_3_sub_val = ws.cell(row=8, column=8).value
    assert student_3_roll_val == "'-"
    assert student_3_sub_val == "Not submitted"


def test_generate_exam_results_excel_empty_results():
    exam = MagicMock(spec=Exam)
    exam.id = "ex_empty"
    exam.title = "Empty Exam"
    exam.subject = "MATHEMATICS"
    exam.maxMarks = 50

    xlsx_bytes = generate_exam_results_excel(exam, [])
    assert xlsx_bytes.startswith(b"PK")

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active
    # Header exists and "No student attempts recorded" message is rendered
    found_empty_msg = False
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            if cell and "No student attempts recorded" in str(cell):
                found_empty_msg = True
                break
    assert found_empty_msg


def test_generate_exam_results_pdf_with_injection_and_unicode():
    exam = MagicMock(spec=Exam)
    exam.id = "ex_202"
    exam.title = "History Quiz <Ancient & Medieval>"
    exam.subject = "HISTORY"
    exam.maxMarks = 50

    results = [
        {
            "rank": 1,
            "attemptId": "att_1",
            "studentName": "Alice & Bob <script>",
            "rollNo": "H-001",
            "marksObtained": 45,
            "maxMarks": 50,
            "percentage": 90.0,
            "status": "SUBMITTED",
            "submittedAt": datetime(2026, 10, 15, 11, 0, tzinfo=UTC),
        },
        {
            "rank": 2,
            "attemptId": "att_2",
            "studentName": "Fatima سلام",  # Arabic / Non-ASCII
            "rollNo": None,
            "marksObtained": 30,
            "maxMarks": 50,
            "percentage": 60.0,
            "status": "IN_PROGRESS",
            "submittedAt": None,
        },
    ]

    pdf_bytes = generate_exam_results_pdf(exam, results)
    assert pdf_bytes.startswith(b"%PDF")
    assert len(pdf_bytes) > 1000


def test_generate_exam_results_pdf_empty_results():
    exam = MagicMock(spec=Exam)
    exam.id = "ex_empty_pdf"
    exam.title = "Empty Exam Results"
    exam.subject = "ENGLISH"
    exam.maxMarks = 100

    pdf_bytes = generate_exam_results_pdf(exam, [])
    assert pdf_bytes.startswith(b"%PDF")
    assert len(pdf_bytes) > 500


def test_generate_exam_results_malformed_exam():
    # Test handling when exam.maxMarks is 0 or None
    exam = MagicMock(spec=Exam)
    exam.id = "ex_malformed"
    exam.title = None
    exam.subject = None
    exam.maxMarks = 0

    results = [
        {
            "rank": 1,
            "attemptId": "att_0",
            "studentName": None,
            "rollNo": None,
            "marksObtained": 0,
            "maxMarks": 0,
            "percentage": 0.0,
            "status": None,
            "submittedAt": None,
        }
    ]

    excel_bytes = generate_exam_results_excel(exam, results)
    assert excel_bytes.startswith(b"PK")

    pdf_bytes = generate_exam_results_pdf(exam, results)
    assert pdf_bytes.startswith(b"%PDF")


@pytest.mark.asyncio
async def test_export_endpoints_authorization(mocker):
    from fastapi import HTTPException

    from app.core.models import Role
    from app.exams.router import (
        export_exam_results_excel_endpoint,
        export_exam_results_pdf_endpoint,
    )
    from app.users.schemas import UserResponse

    student_user = MagicMock(spec=UserResponse)
    student_user.id = "usr_student"
    student_user.role = Role.STUDENT

    mock_exam = MagicMock(spec=Exam)
    mock_exam.id = "ex_secure"
    mock_exam.creatorId = "usr_teacher_other"

    mocker.patch("app.exams.crud.get_exam_by_id", AsyncMock(return_value=mock_exam))
    mocker.patch(
        "app.exams.router.require_exam_manager",
        AsyncMock(side_effect=HTTPException(status_code=403, detail="Not authorized to manage this exam")),
    )

    # Student denied Excel
    with pytest.raises(HTTPException) as exc_excel:
        await export_exam_results_excel_endpoint(exam_id="ex_secure", current_user=student_user)
    assert exc_excel.value.status_code == 403

    # Student denied PDF
    with pytest.raises(HTTPException) as exc_pdf:
        await export_exam_results_pdf_endpoint(exam_id="ex_secure", current_user=student_user)
    assert exc_pdf.value.status_code == 403


@pytest.mark.asyncio
async def test_export_endpoints_success(mocker):
    from app.core.models import Role
    from app.exams.router import (
        export_exam_results_excel_endpoint,
        export_exam_results_pdf_endpoint,
    )
    from app.users.schemas import UserResponse

    teacher_user = MagicMock(spec=UserResponse)
    teacher_user.id = "usr_teacher_1"
    teacher_user.role = Role.TEACHER

    mock_exam = MagicMock(spec=Exam)
    mock_exam.id = "ex_export_ok"
    mock_exam.title = "Maths Test"
    mock_exam.subject = Subject.MATHS
    mock_exam.maxMarks = 100

    mocker.patch("app.exams.crud.get_exam_by_id", AsyncMock(return_value=mock_exam))
    mocker.patch("app.exams.router.require_exam_manager", AsyncMock(return_value=None))
    mocker.patch(
        "app.exams.crud.get_exam_results",
        AsyncMock(return_value=[
            {
                "rank": 1,
                "attemptId": "att_1",
                "studentName": "Alice",
                "rollNo": "A-01",
                "marksObtained": 90,
                "maxMarks": 100,
                "percentage": 90.0,
                "status": "SUBMITTED",
                "submittedAt": datetime(2026, 10, 15, 12, 0, tzinfo=UTC),
            }
        ]),
    )

    # Excel export response
    res_excel = await export_exam_results_excel_endpoint(exam_id="ex_export_ok", current_user=teacher_user)
    assert res_excel.status_code == 200
    assert res_excel.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment; filename=\"exam_results_ex_export_ok.xlsx\"" in res_excel.headers["Content-Disposition"]
    assert res_excel.body.startswith(b"PK")

    # PDF export response
    res_pdf = await export_exam_results_pdf_endpoint(exam_id="ex_export_ok", current_user=teacher_user)
    assert res_pdf.status_code == 200
    assert res_pdf.media_type == "application/pdf"
    assert "attachment; filename=\"exam_results_ex_export_ok.pdf\"" in res_pdf.headers["Content-Disposition"]
    assert res_pdf.body.startswith(b"%PDF")

